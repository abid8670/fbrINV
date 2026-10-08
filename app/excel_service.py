import io
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.models import db, Item, Customer, Invoice, InvoiceItem, CompanySetting

class ExcelService:
    @staticmethod
    def _apply_header_style(ws, headers):
        header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'),
            right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'),
            bottom=Side(style='thin', color='DDDDDD')
        )
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            # Set minimum width
            col_letter = cell.column_letter
            ws.column_dimensions[col_letter].width = max(len(header) + 5, 16)
        ws.row_dimensions[1].height = 28

    @classmethod
    def generate_item_template(cls):
        """Creates a downloadable Excel template for Items Master Data."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Items_Template"

        headers = ["Item Code", "Item Name / Description", "HS Code", "UOM", "Unit Price (PKR)", "Sales Tax Rate (%)"]
        cls._apply_header_style(ws, headers)

        sample_rows = [
            ["ITM-001", "Industrial Steel Screws & Fasteners", "7318.1500", "KG", 450.00, 18.0],
            ["ITM-002", "Cotton Fabric Bleached Grade A", "5208.2100", "MTR", 320.00, 18.0],
            ["ITM-003", "Commercial LED Bulb 18W B22", "8539.5200", "NOS", 275.00, 18.0],
            ["ITM-004", "Heavy Duty Hydraulic Lubricant oil", "2710.1991", "LTR", 1250.00, 18.0]
        ]

        for row_idx, row_data in enumerate(sample_rows, 2):
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = Font(name="Segoe UI", size=10)
                cell.alignment = Alignment(horizontal="left" if isinstance(val, str) else "right")

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def generate_invoice_template(cls):
        """Creates a downloadable Excel template for bulk Invoice Upload."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Invoices_Upload"

        headers = [
            "Invoice Number", "Invoice Date (YYYY-MM-DD)", "Buyer Name", "Buyer Type (Registered/Unregistered)",
            "Buyer NTN", "Buyer CNIC", "Buyer STRN", "Buyer Address", "Buyer Province",
            "Item Code", "Description", "HS Code", "UOM", "Quantity", "Unit Price", "Sales Tax Rate (%)", "Discount"
        ]
        cls._apply_header_style(ws, headers)

        sample_rows = [
            [
                "INV-EXC-001", "2026-10-06", "Prime Builders Corp", "Registered",
                "1234567-8", "", "3277876123456", "Plot 14, I-9 Industrial Area, Islamabad", "Islamabad",
                "ITM-001", "Industrial Steel Screws & Fasteners", "7318.1500", "KG", 50, 450.0, 18.0, 0
            ],
            [
                "INV-EXC-001", "2026-10-06", "Prime Builders Corp", "Registered",
                "1234567-8", "", "3277876123456", "Plot 14, I-9 Industrial Area, Islamabad", "Islamabad",
                "ITM-003", "Commercial LED Bulb 18W B22", "8539.5200", "NOS", 20, 275.0, 18.0, 0
            ],
            [
                "INV-EXC-002", "2026-10-06", "Ahmed General Store", "Unregistered",
                "", "35202-1234567-1", "", "Bazaar Market, Lahore", "Punjab",
                "ITM-002", "Cotton Fabric Bleached Grade A", "5208.2100", "MTR", 100, 320.0, 18.0, 500
            ]
        ]

        for row_idx, row_data in enumerate(sample_rows, 2):
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = Font(name="Segoe UI", size=10)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def import_items(cls, file_stream):
        """Imports items from an uploaded Excel or CSV file."""
        try:
            df = pd.read_excel(file_stream)
        except Exception:
            file_stream.seek(0)
            df = pd.read_csv(file_stream)

        # Normalize column names
        df.columns = [str(c).strip().lower().replace(" ", "_").replace("(", "").replace(")", "").replace("%", "pct").replace("/", "_") for c in df.columns]

        imported_count = 0
        updated_count = 0
        errors = []

        for idx, row in df.iterrows():
            try:
                # Find column values with flexible fallback
                code = str(row.get('item_code') or row.get('code') or f"SKU-{idx+1}").strip()
                name = str(row.get('item_name__description') or row.get('item_name') or row.get('description') or row.get('name') or '').strip()
                if not name:
                    continue

                hs_code = str(row.get('hs_code') or '8471.3010').strip()
                uom = str(row.get('uom') or 'NOS').strip()
                price = float(row.get('unit_price_pkr') or row.get('unit_price') or row.get('price') or 0.0)
                tax_rate = float(row.get('sales_tax_rate_pct') or row.get('sales_tax_rate') or row.get('tax_rate') or 18.0)

                existing = Item.query.filter_by(item_code=code).first()
                if existing:
                    existing.name = name
                    existing.hs_code = hs_code
                    existing.uom = uom
                    existing.unit_price = price
                    existing.sales_tax_rate = tax_rate
                    updated_count += 1
                else:
                    item = Item(
                        item_code=code,
                        name=name,
                        description=name,
                        hs_code=hs_code,
                        uom=uom,
                        unit_price=price,
                        sales_tax_rate=tax_rate
                    )
                    db.session.add(item)
                    imported_count += 1
            except Exception as e:
                errors.append(f"Row {idx+2}: {str(e)}")

        db.session.commit()
        return {
            "success": True,
            "imported": imported_count,
            "updated": updated_count,
            "errors": errors
        }

    @classmethod
    def import_invoices(cls, file_stream, user_id=None):
        """Imports batch invoices from Excel, groups line items, and computes FBR taxes."""
        try:
            df = pd.read_excel(file_stream)
        except Exception:
            file_stream.seek(0)
            df = pd.read_csv(file_stream)

        df.columns = [str(c).strip().lower().replace(" ", "_").replace("(", "").replace(")", "").replace("%", "pct").replace("/", "_").replace("-", "_") for c in df.columns]

        settings = CompanySetting.get_settings()
        created_invoices = 0
        created_ids = []
        errors = []

        # Group by invoice number
        inv_col = 'invoice_number' if 'invoice_number' in df.columns else 'invoice_no'
        if inv_col not in df.columns:
            return {"success": False, "message": "Invoice Number column is required."}

        grouped = df.groupby(inv_col)

        for inv_num, group in grouped:
            try:
                inv_num_clean = str(inv_num).strip()
                first_row = group.iloc[0]

                # Customer resolution
                buyer_name = str(first_row.get('buyer_name') or 'Walk-in Customer').strip()
                buyer_type = str(first_row.get('buyer_type_registered_unregistered') or first_row.get('buyer_type') or 'Registered').strip()
                if buyer_type.lower().startswith('un'):
                    buyer_type = 'Unregistered'
                else:
                    buyer_type = 'Registered'

                buyer_ntn = str(first_row.get('buyer_ntn') or '').strip()
                buyer_cnic = str(first_row.get('buyer_cnic') or '').strip()
                buyer_strn = str(first_row.get('buyer_strn') or '').strip()
                buyer_address = str(first_row.get('buyer_address') or '').strip()
                buyer_province = str(first_row.get('buyer_province') or 'Punjab').strip()

                # Find or create customer
                customer = None
                if buyer_ntn:
                    customer = Customer.query.filter_by(ntn=buyer_ntn).first()
                elif buyer_cnic:
                    customer = Customer.query.filter_by(cnic=buyer_cnic).first()
                if not customer:
                    customer = Customer.query.filter_by(name=buyer_name).first()

                if not customer:
                    customer = Customer(
                        name=buyer_name,
                        buyer_type=buyer_type,
                        ntn=buyer_ntn,
                        cnic=buyer_cnic,
                        strn=buyer_strn,
                        address=buyer_address,
                        province=buyer_province
                    )
                    db.session.add(customer)
                    db.session.flush()

                # Invoice date
                date_val = first_row.get('invoice_date_yyyy_mm_dd') or first_row.get('invoice_date')
                if pd.notna(date_val):
                    if isinstance(date_val, datetime):
                        inv_date = date_val.date()
                    else:
                        try:
                            inv_date = datetime.strptime(str(date_val)[:10], "%Y-%m-%d").date()
                        except Exception:
                            inv_date = datetime.utcnow().date()
                else:
                    inv_date = datetime.utcnow().date()

                # Check if invoice already exists
                invoice = Invoice.query.filter_by(invoice_number=inv_num_clean).first()
                if not invoice:
                    invoice = Invoice(
                        invoice_number=inv_num_clean,
                        invoice_date=inv_date,
                        customer_id=customer.id,
                        status='Draft',
                        created_by_id=user_id
                    )
                    db.session.add(invoice)
                    db.session.flush()
                else:
                    # Clear existing items if reimporting draft
                    InvoiceItem.query.filter_by(invoice_id=invoice.id).delete()

                # Add items
                for _, row in group.iterrows():
                    desc = str(row.get('description') or row.get('item_name') or 'Item Supply').strip()
                    item_code = str(row.get('item_code') or '').strip()
                    hs_code = str(row.get('hs_code') or '8471.3010').strip()
                    uom = str(row.get('uom') or 'NOS').strip()
                    qty = float(row.get('quantity') or row.get('qty') or 1.0)
                    price = float(row.get('unit_price') or row.get('price') or 0.0)
                    tax_rate = float(row.get('sales_tax_rate_pct') or row.get('sales_tax_rate') or settings.default_tax_rate)
                    disc = float(row.get('discount') or 0.0)

                    value_of_supply = round(qty * price, 2)
                    tax_amount = round(value_of_supply * (tax_rate / 100.0), 2)
                    
                    # Further tax for unregistered buyers
                    further_rate = settings.default_further_tax_rate if customer.buyer_type == 'Unregistered' else 0.0
                    further_amount = round(value_of_supply * (further_rate / 100.0), 2)
                    total_amt = round(value_of_supply + tax_amount + further_amount - disc, 2)

                    inv_item = InvoiceItem(
                        invoice_id=invoice.id,
                        item_code=item_code,
                        description=desc,
                        hs_code=hs_code,
                        uom=uom,
                        quantity=qty,
                        unit_price=price,
                        value_of_supply=value_of_supply,
                        sales_tax_rate=tax_rate,
                        sales_tax_amount=tax_amount,
                        further_tax_rate=further_rate,
                        further_tax_amount=further_amount,
                        discount_amount=disc,
                        total_amount=total_amt
                    )
                    db.session.add(inv_item)

                invoice.calculate_totals()
                created_invoices += 1
                created_ids.append(invoice.id)

            except Exception as e:
                errors.append(f"Invoice {inv_num}: {str(e)}")

        db.session.commit()
        return {
            "success": True,
            "invoices_created": created_invoices,
            "invoice_ids": created_ids,
            "errors": errors
        }

    @classmethod
    def export_invoices(cls, invoices):
        """Exports invoices and their FBR statuses into Excel format for FBR reporting & filing."""
        wb = Workbook()
        ws = wb.active
        ws.title = "FBR_Invoices_Report"

        headers = [
            "FBR Status", "FBR Invoice No", "Invoice No", "Date", "Customer Name",
            "Buyer NTN/CNIC", "Buyer Type", "Value of Supply (PKR)", "Sales Tax (PKR)",
            "Further Tax (PKR)", "Discount (PKR)", "Grand Total (PKR)", "FBR Response Code"
        ]
        cls._apply_header_style(ws, headers)

        for row_idx, inv in enumerate(invoices, 2):
            buyer_id = inv.customer.ntn if inv.customer.buyer_type == "Registered" else inv.customer.cnic
            row_data = [
                inv.status,
                inv.fbr_invoice_number or "Not Submitted",
                inv.invoice_number,
                inv.invoice_date.strftime("%Y-%m-%d"),
                inv.customer.name,
                buyer_id or "N/A",
                inv.customer.buyer_type,
                round(inv.total_value_of_supply, 2),
                round(inv.total_sales_tax, 2),
                round(inv.total_further_tax, 2),
                round(inv.total_discount, 2),
                round(inv.grand_total, 2),
                inv.fbr_status_code or ""
            ]
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = Font(name="Segoe UI", size=10)
                if col_idx in [8, 9, 10, 11, 12]:
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = '#,##0.00'

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
