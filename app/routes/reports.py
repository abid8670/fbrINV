import io
from datetime import datetime
from flask import Blueprint, render_template, request, send_file
from flask_login import login_required
from sqlalchemy import func, extract
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.models import Invoice, Customer, CompanySetting
from app.utils import permission_required

reports_bp = Blueprint('reports', __name__, url_prefix='/reports')

@reports_bp.route('/annexure-c')
@login_required
@permission_required('can_export_data')
def annexure_c():
    """Sales Tax Annexure-C Report for Pakistan Sales Tax Monthly Return."""
    year = request.args.get('year', type=int) or datetime.utcnow().year
    month = request.args.get('month', type=int) or datetime.utcnow().month
    status_filter = request.args.get('status', 'FBR_Submitted') # or 'All'

    query = Invoice.query.filter(
        extract('year', Invoice.invoice_date) == year,
        extract('month', Invoice.invoice_date) == month
    )
    if status_filter != 'All':
        query = query.filter_by(status=status_filter)

    invoices = query.order_by(Invoice.invoice_date.asc()).all()

    # Totals breakdown
    reg_value = sum(i.total_value_of_supply for i in invoices if i.customer.buyer_type == 'Registered')
    reg_tax = sum(i.total_sales_tax for i in invoices if i.customer.buyer_type == 'Registered')
    
    unreg_value = sum(i.total_value_of_supply for i in invoices if i.customer.buyer_type == 'Unregistered')
    unreg_tax = sum(i.total_sales_tax for i in invoices if i.customer.buyer_type == 'Unregistered')
    unreg_further = sum(i.total_further_tax for i in invoices if i.customer.buyer_type == 'Unregistered')

    total_value = reg_value + unreg_value
    total_sales_tax = reg_tax + unreg_tax
    total_further_tax = unreg_further
    total_grand = sum(i.grand_total for i in invoices)

    settings = CompanySetting.get_settings()

    return render_template(
        'reports/annexure_c.html',
        invoices=invoices,
        year=year,
        month=month,
        status_filter=status_filter,
        reg_value=reg_value,
        reg_tax=reg_tax,
        unreg_value=unreg_value,
        unreg_tax=unreg_tax,
        unreg_further=unreg_further,
        total_value=total_value,
        total_sales_tax=total_sales_tax,
        total_further_tax=total_further_tax,
        total_grand=total_grand,
        settings=settings
    )

@reports_bp.route('/annexure-c/export')
@login_required
@permission_required('can_export_data')
def export_annexure_c():
    year = request.args.get('year', type=int) or datetime.utcnow().year
    month = request.args.get('month', type=int) or datetime.utcnow().month
    status_filter = request.args.get('status', 'FBR_Submitted')

    query = Invoice.query.filter(
        extract('year', Invoice.invoice_date) == year,
        extract('month', Invoice.invoice_date) == month
    )
    if status_filter != 'All':
        query = query.filter_by(status=status_filter)

    invoices = query.order_by(Invoice.invoice_date.asc()).all()
    settings = CompanySetting.get_settings()

    wb = Workbook()
    ws = wb.active
    ws.title = f"Annex_C_{month}_{year}"

    # Header title
    ws.merge_cells("A1:K1")
    title_cell = ws.cell(row=1, column=1, value=f"{settings.company_name} - Sales Tax Annexure-C ({month:02d}/{year})")
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    headers = [
        "FBR Fiscal No", "Invoice No", "Date", "Buyer Name", "Buyer Type",
        "Buyer NTN/CNIC", "Buyer STRN", "Value of Supply (PKR)", "Sales Tax 18% (PKR)",
        "Further Tax 4% (PKR)", "Total Bill (PKR)"
    ]

    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        col_letter = cell.column_letter
        ws.column_dimensions[col_letter].width = max(len(h) + 4, 15)
    ws.row_dimensions[2].height = 25

    for row_idx, inv in enumerate(invoices, 3):
        buyer_id = inv.customer.ntn if inv.customer.buyer_type == "Registered" else inv.customer.cnic
        row_data = [
            inv.fbr_invoice_number or "Draft",
            inv.invoice_number,
            inv.invoice_date.strftime("%Y-%m-%d"),
            inv.customer.name,
            inv.customer.buyer_type,
            buyer_id or "",
            inv.customer.strn or "",
            round(inv.total_value_of_supply, 2),
            round(inv.total_sales_tax, 2),
            round(inv.total_further_tax, 2),
            round(inv.grand_total, 2)
        ]
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Segoe UI", size=10)
            if col_idx in [8, 9, 10, 11]:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = '#,##0.00'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"FBR_Annexure_C_{year}_{month:02d}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
